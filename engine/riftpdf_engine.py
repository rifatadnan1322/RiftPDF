#!/usr/bin/env python3
"""
RiftPDF engine — the PDF processing backend.

Speaks a tiny JSON-lines protocol on stdout so the Swift app can show live
progress:

    {"type":"progress","value":0.42,"message":"Recompressing images"}
    {"type":"result","...":"..."}         # exactly one, on success
    {"type":"error","message":"...","detail":"..."}

Usage:  riftpdf_engine.py <command> <json-payload-file>
"""

import base64
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import traceback

import pymupdf
import pikepdf

# --------------------------------------------------------------------------
# protocol helpers
# --------------------------------------------------------------------------

# Where emitted messages go. The macOS app spawns this file as a subprocess and
# reads stdout; the Qt app imports it and calls commands on a worker thread, so
# the sink is swappable per thread.
_local = threading.local()


def emit(obj):
    sink = getattr(_local, "sink", None)
    if sink is not None:
        sink(obj)
        return
    sys.stdout.write(json.dumps(obj) + "\n")
    sys.stdout.flush()


def run_command(name, payload, on_progress=None):
    """Run a command in-process and return its result dictionary.

    Raises EngineError with the engine's own message on failure, so callers get
    the same wording users would see from the command line.
    """
    fn = COMMANDS.get(name)
    if fn is None:
        raise EngineError(f"Unknown command '{name}'")

    captured = {}

    def sink(obj):
        kind = obj.get("type")
        if kind == "progress":
            if on_progress:
                on_progress(obj.get("value", 0.0), obj.get("message", ""))
        elif kind == "result":
            captured.update(obj)
        elif kind == "error":
            captured["__error__"] = obj

    previous = getattr(_local, "sink", None)
    _local.sink = sink
    try:
        fn(payload)
    except EngineError:
        raise
    except Exception as exc:
        raise EngineError(str(exc)) from exc
    finally:
        _local.sink = previous

    if "__error__" in captured:
        raise EngineError(captured["__error__"].get("message", "Engine error"))
    captured.pop("type", None)
    return captured


def selftest():
    return {
        "ok": True,
        "pymupdf": pymupdf.version[0],
        "pikepdf": pikepdf.__version__,
        "libreoffice": bool(which("soffice") or which("libreoffice")
                            or os.path.exists("/Applications/LibreOffice.app")),
        "ghostscript": bool(which("gs")),
        "qpdf": bool(which("qpdf")),
        "tesseract": bool(which("tesseract")),
    }


def progress(value, message=""):
    emit({"type": "progress", "value": max(0.0, min(1.0, float(value))), "message": message})


def result(**kwargs):
    payload = {"type": "result"}
    payload.update(kwargs)
    emit(payload)


class EngineError(Exception):
    pass


def human(n):
    n = float(n)
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024


def open_doc(path, password=None):
    try:
        doc = pymupdf.open(path)
    except Exception as e:
        raise EngineError(f"Could not open {os.path.basename(path)}: {e}")
    if doc.needs_pass:
        if not password or not doc.authenticate(password):
            raise EngineError("This PDF is password protected. Supply the open password.")
    return doc


def parse_pages(spec, page_count):
    """'1-3,7,9-' (1-based, inclusive) -> sorted list of 0-based indices."""
    if not spec or str(spec).strip().lower() in ("all", "*", ""):
        return list(range(page_count))
    out = set()
    for chunk in str(spec).replace(" ", "").split(","):
        if not chunk:
            continue
        if "-" in chunk:
            a, _, b = chunk.partition("-")
            start = int(a) - 1 if a else 0
            end = int(b) - 1 if b else page_count - 1
            for i in range(max(0, start), min(page_count - 1, end) + 1):
                out.add(i)
        else:
            i = int(chunk) - 1
            if 0 <= i < page_count:
                out.add(i)
    return sorted(out)


def save_optimised(doc, out_path, linear=True):
    """Save with every safe space saving switch pymupdf offers."""
    kwargs = dict(garbage=4, deflate=True, deflate_images=True,
                  deflate_fonts=True, clean=True, pretty=False)
    try:
        doc.save(out_path, linear=linear, **kwargs)
    except Exception:
        doc.save(out_path, **kwargs)


def which(name):
    return shutil.which(name)


# --------------------------------------------------------------------------
# inspect
# --------------------------------------------------------------------------

def cmd_info(p):
    doc = open_doc(p["input"], p.get("password"))
    md = doc.metadata or {}
    images, fonts = set(), set()
    for pno in range(doc.page_count):
        for img in doc.get_page_images(pno, full=True):
            images.add(img[0])
        for f in doc.get_page_fonts(pno, full=True):
            fonts.add((f[3], f[2], bool(f[1] not in ("n/a", "", None))))
    pages = []
    for i, page in enumerate(doc):
        r = page.rect
        pages.append({"index": i, "width": round(r.width, 2), "height": round(r.height, 2),
                      "rotation": page.rotation})
    size = os.path.getsize(p["input"])
    result(
        pageCount=doc.page_count,
        fileSize=size,
        fileSizeHuman=human(size),
        encrypted=bool(doc.needs_pass) or doc.is_encrypted,
        hasForm=bool(doc.is_form_pdf),
        imageCount=len(images),
        fonts=[{"name": n, "type": t, "embedded": e} for (n, t, e) in sorted(fonts)],
        pages=pages,
        metadata={k: (v or "") for k, v in md.items()},
        hasXMP=bool(doc.xref_xml_metadata() if hasattr(doc, "xref_xml_metadata") else 0),
        toc=[{"level": lvl, "title": title, "page": pg} for lvl, title, pg in doc.get_toc()],
    )


# --------------------------------------------------------------------------
# compress
# --------------------------------------------------------------------------

PRESETS = {
    # name:        (target dpi, jpeg quality, ghostscript preset)
    "light":       (200, 85, "/printer"),
    "balanced":    (150, 72, "/ebook"),
    "aggressive":  (110, 58, "/ebook"),
    "extreme":     (72,  42, "/screen"),
}


def _recompress_images(doc, target_dpi, quality, grayscale=False, report=None):
    """Downsample + re-encode every raster image that is bigger than it needs
    to be for its placement on the page. Returns bytes saved."""
    from PIL import Image

    saved = 0
    seen = set()
    total = doc.page_count or 1
    for pno in range(doc.page_count):
        if report:
            report(pno / total, f"Recompressing images — page {pno + 1} of {total}")
        for info in doc.get_page_images(pno, full=True):
            xref = info[0]
            if xref in seen:
                continue
            seen.add(xref)
            try:
                rects = doc[pno].get_image_rects(xref)
            except Exception:
                rects = []
            try:
                raw = doc.extract_image(xref)
            except Exception:
                continue
            if not raw or not raw.get("image"):
                continue
            original = raw["image"]
            if len(original) < 12 * 1024:      # tiny — not worth the churn
                continue
            try:
                im = Image.open(io.BytesIO(original))
                im.load()
            except Exception:
                continue

            # how many pixels does this image actually need on the page?
            if rects:
                w_pt = max(r.width for r in rects) or 1
                h_pt = max(r.height for r in rects) or 1
                max_w = max(16, int(w_pt / 72.0 * target_dpi))
                max_h = max(16, int(h_pt / 72.0 * target_dpi))
            else:
                max_w, max_h = im.width, im.height

            if im.mode in ("P", "LA", "PA"):
                im = im.convert("RGBA" if "A" in im.mode else "RGB")
            has_alpha = im.mode in ("RGBA", "LA")

            if im.width > max_w or im.height > max_h:
                im.thumbnail((max_w, max_h), Image.LANCZOS)

            if grayscale and not has_alpha:
                im = im.convert("L")

            buf = io.BytesIO()
            if has_alpha:
                im.save(buf, format="PNG", optimize=True)
            else:
                if im.mode not in ("RGB", "L"):
                    im = im.convert("RGB")
                im.save(buf, format="JPEG", quality=quality, optimize=True, progressive=True)
            new = buf.getvalue()

            if len(new) < len(original) * 0.92:
                try:
                    doc.replace_image(xref, stream=new)
                    saved += len(original) - len(new)
                except Exception:
                    pass
    return saved


def _compress_pass(src, out, dpi, quality, grayscale, gs_preset,
                   password=None, subset_fonts=True, strip_metadata=False,
                   flatten=False, use_gs=True, report=None):
    """One complete compression attempt. Returns the resulting file size."""
    def say(frac, msg):
        if report:
            report(frac, msg)

    say(0.02, "Opening document")
    doc = open_doc(src, password)

    _recompress_images(doc, dpi, quality, grayscale,
                       report=lambda f, m: say(0.05 + 0.65 * f, m))

    if subset_fonts:
        say(0.74, "Subsetting fonts")
        try:
            doc.subset_fonts(verbose=False)
        except Exception:
            pass

    if strip_metadata:
        doc.set_metadata({})
        try:
            doc.del_xml_metadata()
        except Exception:
            pass

    if flatten:
        say(0.80, "Flattening annotations")
        try:
            doc.bake()
        except Exception:
            pass

    say(0.84, "Rebuilding file")
    stage1 = out + ".stage1.pdf"
    save_optimised(doc, stage1)
    page_count = doc.page_count
    doc.close()

    best, best_size = stage1, os.path.getsize(stage1)

    gs = which("gs")
    if gs and use_gs:
        say(0.90, "Ghostscript optimisation pass")
        gs_out = out + ".stage2.pdf"
        try:
            subprocess.run(
                [gs, "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.7",
                 f"-dPDFSETTINGS={gs_preset}", "-dNOPAUSE", "-dQUIET", "-dBATCH",
                 "-dDetectDuplicateImages=true", "-dCompressFonts=true",
                 "-dSubsetFonts=true", f"-dColorImageResolution={dpi}",
                 f"-dGrayImageResolution={dpi}", f"-dMonoImageResolution={max(dpi, 300)}",
                 f"-sOutputFile={gs_out}", stage1],
                check=True, capture_output=True, timeout=600)
            if os.path.exists(gs_out) and 0 < os.path.getsize(gs_out) < best_size:
                chk = pymupdf.open(gs_out)
                ok = chk.page_count == page_count
                chk.close()
                if ok:
                    best, best_size = gs_out, os.path.getsize(gs_out)
        except Exception:
            pass

    say(0.97, "Finishing")
    if os.path.exists(out):
        os.remove(out)
    shutil.move(best, out)
    for leftover in (out + ".stage1.pdf", out + ".stage2.pdf"):
        if os.path.exists(leftover):
            os.remove(leftover)
    return os.path.getsize(out)


# quality ladder walked when hunting for a target size, best first
TARGET_LADDER = [
    (300, 92, "/printer", False), (220, 86, "/printer", False),
    (180, 80, "/ebook", False),   (150, 74, "/ebook", False),
    (130, 66, "/ebook", False),   (110, 58, "/ebook", False),
    (96, 50, "/screen", False),   (84, 44, "/screen", False),
    (72, 38, "/screen", False),   (60, 32, "/screen", False),
    (50, 26, "/screen", False),   (42, 22, "/screen", True),
    (36, 18, "/screen", True),    (30, 14, "/screen", True),
]


def cmd_compress_target(p):
    """Compress until the file lands under a size the user asked for.

    Walks a quality ladder by binary search — each rung is a full compression
    attempt — and keeps the highest quality that still fits."""
    src, out = p["input"], p["output"]
    target = int(p["targetBytes"])
    password = p.get("password")
    allow_gray = bool(p.get("allowGrayscale", True))
    before = os.path.getsize(src)

    if target <= 0:
        raise EngineError("Give a target size larger than zero.")

    # Already small enough? Just do a lossless tidy-up.
    if before <= target:
        progress(0.3, "Already under target — optimising losslessly")
        doc = open_doc(src, password)
        save_optimised(doc, out)
        doc.close()
        after = os.path.getsize(out)
        if after > before:
            shutil.copyfile(src, out)
            after = before
        result(before=before, after=after, beforeHuman=human(before),
               afterHuman=human(after), targetBytes=target,
               targetHuman=human(target), hitTarget=True, attempts=0,
               ratio=round(100.0 * (before - after) / before, 1) if before else 0.0,
               settings="lossless", output=out,
               note="The file was already under the target, so nothing was resampled.")
        return

    ladder = [rung for rung in TARGET_LADDER if allow_gray or not rung[3]]
    attempt_dir = tempfile.mkdtemp(prefix="riftpdf-target-")
    attempts = 0
    best_fit = None          # (size, index, path) — highest quality that fits
    smallest = None          # (size, index, path) — fallback if nothing fits

    try:
        lo, hi = 0, len(ladder) - 1
        while lo <= hi:
            mid = (lo + hi) // 2
            dpi, quality, gs_preset, gray = ladder[mid]
            attempts += 1
            candidate = os.path.join(attempt_dir, f"try{attempts}.pdf")
            base = 0.05 + 0.85 * (attempts - 1) / 5.0

            def report(frac, msg, _base=base, _dpi=dpi):
                progress(min(0.95, _base + 0.16 * frac),
                         f"Trying {_dpi} dpi — {msg.lower()}")

            size = _compress_pass(src, candidate, dpi, quality, gray, gs_preset,
                                  password=password,
                                  strip_metadata=bool(p.get("removeMetadata")),
                                  flatten=bool(p.get("flattenAnnotations")),
                                  report=report)

            progress(min(0.95, base + 0.16),
                     f"{human(size)} at {dpi} dpi (target {human(target)})")

            if smallest is None or size < smallest[0]:
                smallest = (size, mid, candidate)

            if size <= target:
                if best_fit is None or mid < best_fit[1]:
                    best_fit = (size, mid, candidate)
                hi = mid - 1          # try for better quality
            else:
                lo = mid + 1          # need to squeeze harder

        chosen = best_fit or smallest
        size, index, path = chosen
        dpi, quality, _, gray = ladder[index]
        shutil.copyfile(path, out)

        after = os.path.getsize(out)
        hit = after <= target
        settings = f"{dpi} dpi, quality {quality}" + (", greyscale" if gray else "")
        note = None
        if not hit:
            note = (f"Could not reach {human(target)} without destroying the pages. "
                    f"{human(after)} is the smallest this document goes while staying readable. "
                    "Removing pages or exporting as images would go further.")

        result(before=before, after=after, beforeHuman=human(before),
               afterHuman=human(after), targetBytes=target, targetHuman=human(target),
               hitTarget=hit, attempts=attempts, settings=settings, output=out,
               ratio=round(100.0 * (before - after) / before, 1) if before else 0.0,
               note=note)
    finally:
        shutil.rmtree(attempt_dir, ignore_errors=True)


def cmd_compress(p):
    src, out = p["input"], p["output"]
    preset = p.get("preset", "balanced")
    dpi, quality, gs_preset = PRESETS.get(preset, PRESETS["balanced"])
    dpi = int(p.get("dpi") or dpi)
    quality = int(p.get("quality") or quality)
    grayscale = bool(p.get("grayscale"))
    before = os.path.getsize(src)

    after = _compress_pass(
        src, out, dpi, quality, grayscale, gs_preset,
        password=p.get("password"),
        subset_fonts=p.get("subsetFonts", True),
        strip_metadata=bool(p.get("removeMetadata")),
        flatten=bool(p.get("flattenAnnotations")),
        use_gs=p.get("useGhostscript", True),
        report=lambda f, m: progress(f, m))
    result(before=before, after=after, beforeHuman=human(before), afterHuman=human(after),
           ratio=round(100.0 * (before - after) / before, 1) if before else 0.0,
           output=out)


# --------------------------------------------------------------------------
# metadata / sanitising
# --------------------------------------------------------------------------

META_KEYS = ["title", "author", "subject", "keywords", "creator", "producer",
             "creationDate", "modDate"]


def cmd_metadata_set(p):
    doc = open_doc(p["input"], p.get("password"))
    md = dict(doc.metadata or {})
    for k, v in (p.get("metadata") or {}).items():
        md[k] = v
    doc.set_metadata(md)
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"])


def cmd_metadata_strip(p):
    """Remove every identifying trace: doc info, XMP, and optionally the
    per-object junk Acrobat leaves behind."""
    src, out = p["input"], p["output"]
    doc = open_doc(src, p.get("password"))
    removed = []

    if doc.metadata and any(doc.metadata.values()):
        removed.append("Document information dictionary")
    doc.set_metadata({})

    try:
        if doc.xref_xml_metadata():
            removed.append("XMP metadata stream")
        doc.del_xml_metadata()
    except Exception:
        pass

    if p.get("removeAnnotationAuthors", True):
        touched = 0
        for page in doc:
            for annot in page.annots() or []:
                inf = annot.info
                if inf.get("title") or inf.get("subject"):
                    inf["title"] = ""
                    inf["subject"] = ""
                    try:
                        annot.set_info(inf)
                        touched += 1
                    except Exception:
                        pass
        if touched:
            removed.append(f"Author names on {touched} annotation(s)")

    if p.get("removeAttachments"):
        n = doc.embfile_count()
        for i in range(n - 1, -1, -1):
            try:
                doc.embfile_del(i)
            except Exception:
                pass
        if n:
            removed.append(f"{n} embedded file(s)")

    doc.save(out, garbage=4, deflate=True, clean=True)
    doc.close()

    # pikepdf handles the trailer bits pymupdf will not touch
    if p.get("resetDocumentID", True):
        try:
            with pikepdf.open(out, allow_overwriting_input=True) as pdf:
                if "/ID" in pdf.trailer:
                    removed.append("Document ID (trailer /ID)")
                meta_removed = False
                with pdf.open_metadata() as meta:
                    for key in list(meta.keys()):
                        del meta[key]
                        meta_removed = True
                if meta_removed and "XMP metadata stream" not in removed:
                    removed.append("Residual XMP properties")
                pdf.save(out, deterministic_id=True, linearize=False)
        except Exception:
            pass

    result(output=out, removed=removed or ["Nothing identifying was found"])


def cmd_sanitize(p):
    """Strip active content — JavaScript, launch actions, embedded files."""
    src, out = p["input"], p["output"]
    findings = []
    with pikepdf.open(src, password=p.get("password") or "") as pdf:
        root = pdf.Root
        if "/Names" in root and "/JavaScript" in root.Names:
            del root.Names["/JavaScript"]
            findings.append("Document-level JavaScript")
        if "/OpenAction" in root:
            del root["/OpenAction"]
            findings.append("Open action")
        if "/AA" in root:
            del root["/AA"]
            findings.append("Additional actions")
        if "/Names" in root and "/EmbeddedFiles" in root.Names:
            del root.Names["/EmbeddedFiles"]
            findings.append("Embedded files")
        for page in pdf.pages:
            if "/AA" in page:
                del page["/AA"]
                findings.append("Page-level actions")
            for annot in page.get("/Annots", []):
                try:
                    if "/A" in annot and annot.A.get("/S") in ("/JavaScript", "/Launch"):
                        del annot["/A"]
                        findings.append("Annotation script/launch action")
                except Exception:
                    pass
        pdf.save(out)
    result(output=out, removed=sorted(set(findings)) or ["No active content found"])


# --------------------------------------------------------------------------
# assembly: merge / split / images
# --------------------------------------------------------------------------

def cmd_merge(p):
    out = pymupdf.open()
    files = p["inputs"]
    toc = []
    for i, item in enumerate(files):
        path = item if isinstance(item, str) else item["path"]
        pw = None if isinstance(item, str) else item.get("password")
        progress(i / max(1, len(files)), f"Adding {os.path.basename(path)}")
        src = open_doc(path, pw)
        if p.get("bookmarkPerFile", True):
            toc.append([1, os.path.splitext(os.path.basename(path))[0], out.page_count + 1])
        out.insert_pdf(src)
        src.close()
    if toc:
        out.set_toc(toc)
    progress(0.95, "Writing")
    save_optimised(out, p["output"])
    result(output=p["output"], pageCount=out.page_count)


def cmd_split(p):
    doc = open_doc(p["input"], p.get("password"))
    mode = p.get("mode", "ranges")
    outdir = p["outputDir"]
    os.makedirs(outdir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(p["input"]))[0]
    written = []

    if mode == "every":
        n = int(p.get("every", 1))
        groups = [list(range(i, min(i + n, doc.page_count)))
                  for i in range(0, doc.page_count, n)]
    elif mode == "ranges":
        groups = [parse_pages(r, doc.page_count) for r in p.get("ranges", [])]
    else:  # "each"
        groups = [[i] for i in range(doc.page_count)]

    for i, grp in enumerate(groups):
        if not grp:
            continue
        progress(i / max(1, len(groups)), f"Writing part {i + 1}")
        part = pymupdf.open()
        for pno in grp:
            part.insert_pdf(doc, from_page=pno, to_page=pno)
        label = f"{grp[0] + 1}" if len(grp) == 1 else f"{grp[0] + 1}-{grp[-1] + 1}"
        path = os.path.join(outdir, f"{stem} {label}.pdf")
        save_optimised(part, path, linear=False)
        part.close()
        written.append(path)
    result(files=written, outputDir=outdir)


PAGE_SIZES = {
    "letter": (612, 792), "legal": (612, 1008), "a4": (595, 842),
    "a3": (842, 1191), "tabloid": (792, 1224),
}


def cmd_images_to_pdf(p):
    from PIL import Image, ImageOps

    files = p["inputs"]
    fit = p.get("fit", "fit")            # fit | fill-page | actual
    size_name = (p.get("pageSize") or "auto").lower()
    margin = float(p.get("margin", 0))
    landscape_auto = p.get("autoOrient", True)
    out = pymupdf.open()

    for i, path in enumerate(files):
        progress(i / max(1, len(files)), f"Placing {os.path.basename(path)}")
        try:
            im = Image.open(path)
            im = ImageOps.exif_transpose(im)
            if im.mode in ("RGBA", "LA", "P"):
                bg = Image.new("RGB", im.size, "white")
                im = im.convert("RGBA")
                bg.paste(im, mask=im.split()[-1])
                im = bg
            elif im.mode != "RGB":
                im = im.convert("RGB")
        except Exception as e:
            raise EngineError(f"{os.path.basename(path)} is not a readable image: {e}")

        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=int(p.get("quality", 88)), optimize=True)
        data = buf.getvalue()

        iw, ih = im.size
        if size_name == "auto":
            # one page exactly the size of the image at 72dpi-equivalent
            dpi = float(p.get("imageDPI", 150)) or 150
            pw, ph = iw / dpi * 72, ih / dpi * 72
        else:
            pw, ph = PAGE_SIZES.get(size_name, PAGE_SIZES["letter"])
            if landscape_auto and iw > ih:
                pw, ph = ph, pw

        page = out.new_page(width=pw + 2 * margin, height=ph + 2 * margin)
        box = pymupdf.Rect(margin, margin, pw + margin, ph + margin)
        if fit == "actual":
            page.insert_image(box, stream=data, keep_proportion=True)
        else:
            page.insert_image(box, stream=data, keep_proportion=(fit != "fill-page"))

    progress(0.95, "Writing")
    save_optimised(out, p["output"])
    result(output=p["output"], pageCount=out.page_count)


def cmd_pdf_to_images(p):
    doc = open_doc(p["input"], p.get("password"))
    outdir = p["outputDir"]
    os.makedirs(outdir, exist_ok=True)
    dpi = int(p.get("dpi", 200))
    fmt = p.get("format", "png").lower()
    stem = os.path.splitext(os.path.basename(p["input"]))[0]
    pages = parse_pages(p.get("pages"), doc.page_count)
    written = []
    for i, pno in enumerate(pages):
        progress(i / max(1, len(pages)), f"Rendering page {pno + 1}")
        pix = doc[pno].get_pixmap(dpi=dpi, alpha=False)
        path = os.path.join(outdir, f"{stem} page {pno + 1}.{ 'jpg' if fmt in ('jpg','jpeg') else fmt }")
        if fmt in ("jpg", "jpeg"):
            pix.pil_save(path, format="JPEG", quality=int(p.get("quality", 90)))
        else:
            pix.save(path)
        written.append(path)
    result(files=written, outputDir=outdir)


def cmd_extract_images(p):
    doc = open_doc(p["input"], p.get("password"))
    outdir = p["outputDir"]
    os.makedirs(outdir, exist_ok=True)
    seen, written = set(), []
    minpx = int(p.get("minPixels", 64))
    for pno in range(doc.page_count):
        progress(pno / max(1, doc.page_count), f"Scanning page {pno + 1}")
        for info in doc.get_page_images(pno, full=True):
            xref = info[0]
            if xref in seen:
                continue
            seen.add(xref)
            try:
                raw = doc.extract_image(xref)
            except Exception:
                continue
            if not raw:
                continue
            if raw.get("width", 0) < minpx or raw.get("height", 0) < minpx:
                continue
            path = os.path.join(outdir, f"image {len(written) + 1}.{raw['ext']}")
            with open(path, "wb") as fh:
                fh.write(raw["image"])
            written.append(path)
    result(files=written, outputDir=outdir, count=len(written))


# --------------------------------------------------------------------------
# conversion
# --------------------------------------------------------------------------

def cmd_pdf_to_word(p):
    from pdf2docx import Converter
    progress(0.05, "Analysing layout")
    cv = Converter(p["input"], password=p.get("password") or None)
    try:
        pages = p.get("pages")
        kwargs = {}
        if pages:
            idx = parse_pages(pages, 10 ** 6)
            kwargs["pages"] = idx
        cv.convert(p["output"], **kwargs)
    finally:
        cv.close()
    progress(1.0, "Done")
    result(output=p["output"], size=os.path.getsize(p["output"]))


def cmd_pdf_to_text(p):
    doc = open_doc(p["input"], p.get("password"))
    pages = parse_pages(p.get("pages"), doc.page_count)
    mode = p.get("mode", "text")     # text | markdownish
    chunks = []
    for i, pno in enumerate(pages):
        progress(i / max(1, len(pages)), f"Extracting page {pno + 1}")
        chunks.append(doc[pno].get_text("text"))
    body = ("\n\n" + "-" * 40 + "\n\n").join(chunks) if p.get("pageBreaks", True) else "\n".join(chunks)
    with open(p["output"], "w", encoding="utf-8") as fh:
        fh.write(body)
    result(output=p["output"], characters=len(body))


def cmd_office_to_pdf(p):
    """High fidelity path — only available when LibreOffice is installed.
    The app falls back to its native converter when this reports unavailable."""
    soffice = which("soffice") or which("libreoffice")
    for candidate in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",):
        if not soffice and os.path.exists(candidate):
            soffice = candidate
    if not soffice:
        raise EngineError("LIBREOFFICE_UNAVAILABLE")
    outdir = os.path.dirname(p["output"]) or "."
    progress(0.2, "Converting with LibreOffice")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--norestore", "--convert-to", "pdf",
                        "--outdir", tmp, p["input"]],
                       check=True, capture_output=True, timeout=300)
        produced = [f for f in os.listdir(tmp) if f.lower().endswith(".pdf")]
        if not produced:
            raise EngineError("LibreOffice produced no output")
        shutil.move(os.path.join(tmp, produced[0]), p["output"])
    result(output=p["output"])


# --------------------------------------------------------------------------
# stamping: watermark, page numbers, headers
# --------------------------------------------------------------------------

def _anchor_point(rect, position, margin, w, h):
    x0, y0, x1, y1 = rect.x0, rect.y0, rect.x1, rect.y1
    horiz = {"left": x0 + margin, "center": (x0 + x1 - w) / 2, "right": x1 - margin - w}
    vert = {"top": y0 + margin, "middle": (y0 + y1 - h) / 2, "bottom": y1 - margin - h}
    v, _, hgt = position.partition("-")
    return horiz.get(hgt or "center", horiz["center"]), vert.get(v, vert["bottom"])


FONT_MAP = {
    ("helvetica", False, False): "helv", ("helvetica", True, False): "hebo",
    ("helvetica", False, True): "heit", ("helvetica", True, True): "hebi",
    ("times", False, False): "tiro", ("times", True, False): "tibo",
    ("times", False, True): "tiit", ("times", True, True): "tibi",
    ("courier", False, False): "cour", ("courier", True, False): "cobo",
    ("courier", False, True): "coit", ("courier", True, True): "cobi",
}


def pick_font(family="helvetica", bold=False, italic=False):
    fam = (family or "helvetica").lower()
    for key in ("times", "courier", "helvetica"):
        if key in fam:
            fam = key
            break
    else:
        fam = "helvetica"
    return FONT_MAP.get((fam, bool(bold), bool(italic)), "helv")


def cmd_watermark(p):
    doc = open_doc(p["input"], p.get("password"))
    pages = parse_pages(p.get("pages"), doc.page_count)
    opacity = float(p.get("opacity", 0.18))
    rotate = float(p.get("rotate", 45))
    color = p.get("color", [0.55, 0.55, 0.6])
    on_top = bool(p.get("onTop", True))
    kind = p.get("kind", "text")

    for i, pno in enumerate(pages):
        progress(i / max(1, len(pages)), f"Stamping page {pno + 1}")
        page = doc[pno]
        rect = page.rect
        if kind == "image":
            scale = float(p.get("scale", 0.5))
            w = rect.width * scale
            h = w
            box = pymupdf.Rect(0, 0, w, h)
            box = box + ((rect.width - w) / 2, (rect.height - h) / 2,
                         (rect.width - w) / 2, (rect.height - h) / 2)
            page.insert_image(box, filename=p["imagePath"], overlay=on_top,
                              rotate=int(rotate) // 90 * 90, keep_proportion=True)
        else:
            text = p.get("text", "DRAFT")
            size = float(p.get("fontSize", 0) or max(18, rect.width / max(6, len(text)) * 1.6))
            font = pick_font(p.get("font"), p.get("bold", True), p.get("italic", False))
            tw = pymupdf.get_text_length(text, fontname=font, fontsize=size)
            center = pymupdf.Point(rect.width / 2, rect.height / 2)
            box = pymupdf.Rect(center.x - tw / 2 - 8, center.y - size,
                               center.x + tw / 2 + 8, center.y + size)
            page.insert_textbox(box, text, fontname=font, fontsize=size,
                                color=color, align=pymupdf.TEXT_ALIGN_CENTER,
                                rotate=int(rotate) % 360 // 90 * 90,
                                fill_opacity=opacity, stroke_opacity=opacity,
                                overlay=on_top)
    progress(0.95, "Writing")
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"], pagesStamped=len(pages))


def cmd_page_numbers(p):
    doc = open_doc(p["input"], p.get("password"))
    pages = parse_pages(p.get("pages"), doc.page_count)
    fmt = p.get("format", "{n}")
    start = int(p.get("startAt", 1))
    size = float(p.get("fontSize", 10))
    margin = float(p.get("margin", 32))
    position = p.get("position", "bottom-center")
    color = p.get("color", [0.25, 0.25, 0.28])
    font = pick_font(p.get("font"), p.get("bold"), p.get("italic"))
    total = len(pages)

    for i, pno in enumerate(pages):
        n = start + i
        label = (fmt.replace("{n}", str(n))
                    .replace("{total}", str(total))
                    .replace("{page}", str(pno + 1))
                    .replace("{filename}", os.path.splitext(os.path.basename(p["input"]))[0]))
        page = doc[pno]
        w = pymupdf.get_text_length(label, fontname=font, fontsize=size)
        x, y = _anchor_point(page.rect, position, margin, w, size * 1.3)
        page.insert_textbox(pymupdf.Rect(x, y, x + w + 2, y + size * 1.4), label,
                            fontname=font, fontsize=size, color=color, overlay=True)
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"], pagesNumbered=len(pages))


# --------------------------------------------------------------------------
# redaction — genuine removal, not a black rectangle
# --------------------------------------------------------------------------

def cmd_redact(p):
    doc = open_doc(p["input"], p.get("password"))
    by_page = {}
    for item in p["areas"]:
        by_page.setdefault(int(item["page"]), []).append(item)
    fill = p.get("fill", [0, 0, 0])

    for i, (pno, items) in enumerate(sorted(by_page.items())):
        progress(i / max(1, len(by_page)), f"Redacting page {pno + 1}")
        page = doc[pno]
        for item in items:
            r = pymupdf.Rect(*item["rect"])
            page.add_redact_annot(r, fill=fill if fill else None,
                                  text=item.get("overlayText") or None,
                                  fontsize=float(item.get("fontSize", 9)))
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS
                              if p.get("scrubImages", True) else pymupdf.PDF_REDACT_IMAGE_NONE)
    if p.get("stripMetadata", True):
        doc.set_metadata({})
        try:
            doc.del_xml_metadata()
        except Exception:
            pass
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"], areas=len(p["areas"]))


def cmd_redact_search(p):
    """Find every occurrence of the given terms and redact them outright."""
    doc = open_doc(p["input"], p.get("password"))
    terms = p["terms"]
    hits = 0
    for pno in range(doc.page_count):
        progress(pno / max(1, doc.page_count), f"Searching page {pno + 1}")
        page = doc[pno]
        found = False
        for term in terms:
            for rect in page.search_for(term, quads=False):
                page.add_redact_annot(rect, fill=p.get("fill", [0, 0, 0]))
                hits += 1
                found = True
        if found:
            page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE)
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"], occurrences=hits)


# --------------------------------------------------------------------------
# real text editing
# --------------------------------------------------------------------------

def cmd_text_spans(p):
    """Report every editable text run on a page with position and styling."""
    doc = open_doc(p["input"], p.get("password"))
    pno = int(p.get("page", 0))
    page = doc[pno]
    data = page.get_text("dict")
    blocks = []
    for b in data.get("blocks", []):
        if b.get("type") != 0:
            continue
        lines = []
        for ln in b.get("lines", []):
            spans = []
            for sp in ln.get("spans", []):
                if not sp.get("text", "").strip():
                    continue
                col = sp.get("color", 0)
                spans.append({
                    "text": sp["text"],
                    "bbox": [round(v, 2) for v in sp["bbox"]],
                    "font": sp.get("font", ""),
                    "size": round(sp.get("size", 11), 2),
                    "color": [((col >> 16) & 255) / 255.0, ((col >> 8) & 255) / 255.0, (col & 255) / 255.0],
                    "bold": bool(sp.get("flags", 0) & 2 ** 4),
                    "italic": bool(sp.get("flags", 0) & 2 ** 1),
                })
            if spans:
                lines.append({"bbox": [round(v, 2) for v in ln["bbox"]], "spans": spans})
        if lines:
            blocks.append({"bbox": [round(v, 2) for v in b["bbox"]], "lines": lines})
    result(page=pno, blocks=blocks,
           pageWidth=round(page.rect.width, 2), pageHeight=round(page.rect.height, 2))


def cmd_text_edit(p):
    """Replace text in place: erase the original glyphs, lay the new run down
    with matching styling."""
    doc = open_doc(p["input"], p.get("password"))
    edits = p["edits"]
    by_page = {}
    for e in edits:
        by_page.setdefault(int(e["page"]), []).append(e)

    for i, (pno, items) in enumerate(sorted(by_page.items())):
        progress(i / max(1, len(by_page)), f"Rewriting page {pno + 1}")
        page = doc[pno]
        # 1. erase originals without disturbing artwork underneath
        for e in items:
            page.add_redact_annot(pymupdf.Rect(*e["bbox"]), fill=None)
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                              graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
        # 2. lay down replacements
        for e in items:
            text = e.get("text", "")
            if not text:
                continue
            x0, y0, x1, y1 = e["bbox"]
            size = float(e.get("size", 11))
            font = pick_font(e.get("font", ""), e.get("bold"), e.get("italic"))
            color = e.get("color", [0, 0, 0])
            align = {"left": 0, "center": 1, "right": 2, "justify": 3}.get(e.get("align", "left"), 0)
            box = pymupdf.Rect(x0 - 1, y0 - 1, max(x1 + 2, x0 + 12), y1 + max(3, size * 0.4))
            # shrink to fit rather than silently dropping the overflow
            for attempt in range(14):
                rc = page.insert_textbox(box, text, fontname=font, fontsize=size,
                                         color=color, align=align, overlay=True)
                if rc >= 0:
                    break
                size *= 0.94
                if size < 3:
                    box.y1 += 6
                    size = float(e.get("size", 11))
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"], edits=len(edits))


def cmd_add_text(p):
    """Drop brand new text onto a page (the 'add content' tool)."""
    doc = open_doc(p["input"], p.get("password"))
    for item in p["items"]:
        page = doc[int(item["page"])]
        font = pick_font(item.get("font"), item.get("bold"), item.get("italic"))
        align = {"left": 0, "center": 1, "right": 2, "justify": 3}.get(item.get("align", "left"), 0)
        page.insert_textbox(pymupdf.Rect(*item["rect"]), item["text"],
                            fontname=font, fontsize=float(item.get("size", 12)),
                            color=item.get("color", [0, 0, 0]), align=align,
                            fill_opacity=float(item.get("opacity", 1)), overlay=True)
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"])


def cmd_place_images(p):
    """Stamp images (signatures, logos, photos) into the page content itself so
    they survive in every other reader."""
    doc = open_doc(p["input"], p.get("password"))
    for item in p["items"]:
        page = doc[int(item["page"])]
        rect = pymupdf.Rect(*item["rect"])
        kwargs = dict(keep_proportion=bool(item.get("keepProportion", True)), overlay=True)
        if item.get("rotate"):
            kwargs["rotate"] = int(item["rotate"]) // 90 * 90
        if item.get("data"):
            page.insert_image(rect, stream=base64.b64decode(item["data"]), **kwargs)
        else:
            page.insert_image(rect, filename=item["path"], **kwargs)
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"], placed=len(p["items"]))


def cmd_ocr_layer(p):
    """Take word boxes recognised by the app's Vision pass and weld them into
    the PDF as a selectable, searchable, invisible text layer."""
    doc = open_doc(p["input"], p.get("password"))
    by_page = {}
    for w in p["words"]:
        by_page.setdefault(int(w["page"]), []).append(w)
    for i, (pno, words) in enumerate(sorted(by_page.items())):
        progress(i / max(1, len(by_page)), f"Adding text layer to page {pno + 1}")
        page = doc[pno]
        for w in words:
            x0, y0, x1, y1 = w["rect"]
            h = max(1.0, y1 - y0)
            size = h * 0.86
            length = pymupdf.get_text_length(w["text"], fontname="helv", fontsize=size) or 1
            width = max(1.0, x1 - x0)
            page.insert_text(pymupdf.Point(x0, y1 - h * 0.18), w["text"],
                             fontname="helv", fontsize=size,
                             render_mode=3,                    # invisible
                             morph=(pymupdf.Point(x0, y1 - h * 0.18),
                                    pymupdf.Matrix(width / length, 0, 0, 1, 0, 0)))
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"], words=len(p["words"]))


# --------------------------------------------------------------------------
# security
# --------------------------------------------------------------------------

def cmd_encrypt(p):
    doc = open_doc(p["input"], p.get("password"))
    perms = p.get("permissions", {})
    bits = 0
    table = {
        "print": pymupdf.PDF_PERM_PRINT,
        "modify": pymupdf.PDF_PERM_MODIFY,
        "copy": pymupdf.PDF_PERM_COPY,
        "annotate": pymupdf.PDF_PERM_ANNOTATE,
        "fillForms": pymupdf.PDF_PERM_FORM,
        "accessibility": pymupdf.PDF_PERM_ACCESSIBILITY,
        "assemble": pymupdf.PDF_PERM_ASSEMBLE,
        "printHighRes": pymupdf.PDF_PERM_PRINT_HQ,
    }
    for key, flag in table.items():
        if perms.get(key, True):
            bits |= flag
    doc.save(p["output"], encryption=pymupdf.PDF_ENCRYPT_AES_256,
             owner_pw=p.get("ownerPassword") or p.get("userPassword") or None,
             user_pw=p.get("userPassword") or None,
             permissions=bits, garbage=3, deflate=True)
    result(output=p["output"], encryption="AES-256")


def cmd_decrypt(p):
    doc = open_doc(p["input"], p.get("password"))
    doc.save(p["output"], encryption=pymupdf.PDF_ENCRYPT_NONE, garbage=3, deflate=True)
    result(output=p["output"])


# --------------------------------------------------------------------------
# structure: pages, bookmarks, attachments, forms
# --------------------------------------------------------------------------

def cmd_page_ops(p):
    """Batch page surgery in one pass: reorder, delete, rotate, crop, scale."""
    doc = open_doc(p["input"], p.get("password"))

    for op in p.get("operations", []):
        kind = op.get("op")
        if kind == "delete":
            for pno in sorted(parse_pages(op.get("pages"), doc.page_count), reverse=True):
                doc.delete_page(pno)
        elif kind == "reorder":
            doc.select(op["order"])
        elif kind == "rotate":
            for pno in parse_pages(op.get("pages"), doc.page_count):
                page = doc[pno]
                page.set_rotation((page.rotation + int(op.get("degrees", 90))) % 360)
        elif kind == "crop":
            box = op["rect"]
            for pno in parse_pages(op.get("pages"), doc.page_count):
                page = doc[pno]
                r = page.rect
                doc[pno].set_cropbox(pymupdf.Rect(
                    r.x0 + box[0], r.y0 + box[1], r.x1 - box[2], r.y1 - box[3]))
        elif kind == "resize":
            target = PAGE_SIZES.get((op.get("pageSize") or "letter").lower(), PAGE_SIZES["letter"])
            new = pymupdf.open()
            for pno in range(doc.page_count):
                src = doc[pno]
                w, h = target
                if src.rect.width > src.rect.height:
                    w, h = h, w
                page = new.new_page(width=w, height=h)
                page.show_pdf_page(page.rect, doc, pno)
            doc.close()
            doc = new
        elif kind == "insertBlank":
            at = int(op.get("at", doc.page_count))
            w, h = PAGE_SIZES.get((op.get("pageSize") or "letter").lower(), PAGE_SIZES["letter"])
            if doc.page_count and op.get("matchPrevious", True):
                ref = doc[max(0, min(at, doc.page_count) - 1)].rect
                w, h = ref.width, ref.height
            doc.new_page(pno=at, width=w, height=h)
        elif kind == "insertPDF":
            other = open_doc(op["path"], op.get("password"))
            doc.insert_pdf(other, start_at=int(op.get("at", doc.page_count)))
            other.close()

    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"], pageCount=doc.page_count)


def cmd_bookmarks_set(p):
    doc = open_doc(p["input"], p.get("password"))
    doc.set_toc([[int(b["level"]), b["title"], int(b["page"])] for b in p["bookmarks"]])
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"], count=len(p["bookmarks"]))


def cmd_attachments(p):
    doc = open_doc(p["input"], p.get("password"))
    action = p.get("action", "list")
    if action == "list":
        items = []
        for i in range(doc.embfile_count()):
            info = doc.embfile_info(i)
            items.append({"index": i, "name": info.get("filename", ""),
                          "size": info.get("size", 0), "description": info.get("desc", "")})
        result(attachments=items)
        return
    if action == "add":
        for path in p["files"]:
            with open(path, "rb") as fh:
                doc.embfile_add(os.path.basename(path), fh.read(),
                                filename=os.path.basename(path))
    elif action == "extract":
        outdir = p["outputDir"]
        os.makedirs(outdir, exist_ok=True)
        written = []
        for i in range(doc.embfile_count()):
            info = doc.embfile_info(i)
            dest = os.path.join(outdir, info.get("filename") or f"attachment {i + 1}")
            with open(dest, "wb") as fh:
                fh.write(doc.embfile_get(i))
            written.append(dest)
        result(files=written, outputDir=outdir)
        return
    elif action == "remove":
        for i in sorted(p.get("indices", []), reverse=True):
            doc.embfile_del(i)
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"])


def cmd_form_fields(p):
    doc = open_doc(p["input"], p.get("password"))
    fields = []
    for pno in range(doc.page_count):
        for w in doc[pno].widgets():
            fields.append({
                "page": pno, "name": w.field_name or "", "value": str(w.field_value or ""),
                "type": w.field_type_string, "rect": [round(v, 2) for v in w.rect],
                "options": list(w.choice_values or []) if w.choice_values else [],
                "readOnly": bool(w.field_flags & 1),
            })
    result(fields=fields, hasForm=bool(doc.is_form_pdf))


def cmd_form_fill(p):
    doc = open_doc(p["input"], p.get("password"))
    values = {v["name"]: v["value"] for v in p["values"]}
    filled = 0
    for pno in range(doc.page_count):
        for w in doc[pno].widgets():
            if w.field_name in values:
                val = values[w.field_name]
                if w.field_type == pymupdf.PDF_WIDGET_TYPE_CHECKBOX:
                    w.field_value = bool(val) and str(val).lower() not in ("false", "0", "off", "")
                else:
                    w.field_value = str(val)
                w.update()
                filled += 1
    if p.get("flatten"):
        doc.bake(annots=False, widgets=True)
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"], filled=filled)


def cmd_flatten(p):
    doc = open_doc(p["input"], p.get("password"))
    doc.bake(annots=bool(p.get("annotations", True)), widgets=bool(p.get("widgets", True)))
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"])


def cmd_repair(p):
    """Rebuild a damaged file — cross reference tables, broken streams, the lot."""
    src, out = p["input"], p["output"]
    notes = []
    try:
        with pikepdf.open(src, password=p.get("password") or "", allow_overwriting_input=False) as pdf:
            pdf.save(out, fix_metadata_version=True, recompress_flate=True)
        notes.append("Rebuilt cross-reference table")
    except Exception as e:
        notes.append(f"pikepdf could not open it ({e}); falling back to MuPDF")
        doc = pymupdf.open(src)
        doc.save(out, garbage=4, clean=True, deflate=True)
        notes.append("Reconstructed with MuPDF")
    chk = pymupdf.open(out)
    result(output=out, pageCount=chk.page_count, notes=notes)


def cmd_linearize(p):
    """Optimise for fast web view — first page streams before the rest."""
    q = which("qpdf")
    if q:
        subprocess.run([q, "--linearize", "--object-streams=generate",
                        p["input"], p["output"]], check=False, capture_output=True)
        if os.path.exists(p["output"]) and os.path.getsize(p["output"]) > 0:
            result(output=p["output"], method="qpdf")
            return
    doc = open_doc(p["input"], p.get("password"))
    save_optimised(doc, p["output"], linear=True)
    result(output=p["output"], method="mupdf")


def cmd_compare(p):
    """Word-level diff between two PDFs, page by page."""
    import difflib
    a = open_doc(p["inputA"], p.get("passwordA"))
    b = open_doc(p["inputB"], p.get("passwordB"))
    pages = []
    for i in range(max(a.page_count, b.page_count)):
        ta = a[i].get_text("text").split() if i < a.page_count else []
        tb = b[i].get_text("text").split() if i < b.page_count else []
        sm = difflib.SequenceMatcher(None, ta, tb)
        changes = []
        for tag, i1, i2, j1, j2 in sm.get_opcodes():
            if tag == "equal":
                continue
            changes.append({"kind": tag,
                            "before": " ".join(ta[i1:i2])[:400],
                            "after": " ".join(tb[j1:j2])[:400]})
        pages.append({"page": i, "similarity": round(sm.ratio(), 3), "changes": changes})
    result(pages=pages,
           identical=all(pg["similarity"] == 1.0 and not pg["changes"] for pg in pages))


# --------------------------------------------------------------------------
# accessibility
# --------------------------------------------------------------------------

def cmd_accessibility_check(p):
    """Audit the document the way Acrobat's accessibility checker does, and say
    plainly which problems RiftPDF can fix on its own."""
    src = p["input"]
    doc = open_doc(src, p.get("password"))
    issues = []

    def add(severity, title, detail, fixable, rule):
        issues.append({"severity": severity, "title": title, "detail": detail,
                       "fixable": fixable, "rule": rule})

    md = doc.metadata or {}

    # --- document level -------------------------------------------------
    if not (md.get("title") or "").strip():
        add("error", "No document title",
            "Screen readers announce the file name instead of a real title.",
            True, "title")

    lang = None
    try:
        with pikepdf.open(src, password=p.get("password") or "") as pdf:
            lang = str(pdf.Root.get("/Lang", "")) or None
            has_struct = "/StructTreeRoot" in pdf.Root
            marked = False
            if "/MarkInfo" in pdf.Root:
                marked = bool(pdf.Root.MarkInfo.get("/Marked", False))
            shows_title = False
            if "/ViewerPreferences" in pdf.Root:
                shows_title = bool(pdf.Root.ViewerPreferences.get("/DisplayDocTitle", False))
    except Exception:
        has_struct, marked, shows_title = False, False, False

    if not lang:
        add("error", "No document language",
            "Without a language, a screen reader may read the text with the wrong pronunciation.",
            True, "lang")

    if not has_struct:
        add("error", "Not tagged",
            "The document has no structure tree, so assistive technology cannot tell "
            "headings from body text, or follow the reading order. This needs a real "
            "tagging pass — RiftPDF cannot invent structure that was never authored.",
            False, "tagged")
    elif not marked:
        add("warning", "Structure present but not marked as tagged",
            "A structure tree exists but /MarkInfo does not declare the file tagged.",
            True, "marked")

    if not shows_title:
        add("warning", "Window shows the file name, not the title",
            "Turning on DisplayDocTitle makes readers announce the real title.",
            True, "displayTitle")

    if doc.is_encrypted:
        try:
            perms = doc.permissions
            if not (perms & pymupdf.PDF_PERM_ACCESSIBILITY):
                add("error", "Screen reader access is blocked by permissions",
                    "The security settings forbid content extraction for accessibility.",
                    True, "permissions")
        except Exception:
            pass

    # --- page level -----------------------------------------------------
    scanned_pages, image_count, alt_missing, tiny_text = [], 0, 0, 0
    for pno in range(doc.page_count):
        progress(pno / max(1, doc.page_count), f"Checking page {pno + 1}")
        page = doc[pno]
        text = page.get_text("text").strip()
        images = doc.get_page_images(pno, full=True)
        image_count += len(images)
        if images and len(text) < 24:
            scanned_pages.append(pno + 1)
        # every image needs a description; without a struct tree none have one
        if not has_struct:
            alt_missing += len(images)
        for block in page.get_text("dict").get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    if span.get("size", 12) < 6 and span.get("text", "").strip():
                        tiny_text += 1

    if scanned_pages:
        shown = ", ".join(str(n) for n in scanned_pages[:8])
        more = "" if len(scanned_pages) <= 8 else f" and {len(scanned_pages) - 8} more"
        add("error", f"{len(scanned_pages)} page(s) are images with no text",
            f"Pages {shown}{more} look like scans. Nothing on them can be read aloud, "
            "searched or selected until text is recognised.",
            True, "ocr")

    if alt_missing:
        add("error", f"{alt_missing} image(s) have no alternative text",
            "Images need descriptions for screen reader users. This requires a tag "
            "tree, so it goes hand in hand with tagging the document.",
            False, "alt")

    if tiny_text:
        add("warning", f"{tiny_text} run(s) of text below 6 pt",
            "Very small text is hard to read and often unreadable when printed.",
            False, "tinyText")

    if doc.is_form_pdf:
        unlabelled = 0
        for pno in range(doc.page_count):
            for w in doc[pno].widgets():
                if not (w.field_name or "").strip():
                    unlabelled += 1
        if unlabelled:
            add("error", f"{unlabelled} form field(s) have no name",
                "Screen readers read the field name as its label.", False, "fieldNames")

    score = 100
    for issue in issues:
        score -= 18 if issue["severity"] == "error" else 7
    score = max(0, score)

    result(issues=issues, score=score,
           errors=sum(1 for i in issues if i["severity"] == "error"),
           warnings=sum(1 for i in issues if i["severity"] == "warning"),
           pageCount=doc.page_count, imageCount=image_count,
           fixable=[i["rule"] for i in issues if i["fixable"]],
           tagged=has_struct)


def cmd_accessibility_fix(p):
    """Apply the repairs that can be made honestly, and report exactly which."""
    src, out = p["input"], p["output"]
    rules = set(p.get("rules") or [])
    applied, skipped = [], []

    doc = open_doc(src, p.get("password"))
    md = dict(doc.metadata or {})

    if "title" in rules:
        title = (p.get("title") or "").strip()
        if title:
            md["title"] = title
            doc.set_metadata(md)
            applied.append(f"Set document title to “{title}”")
        else:
            skipped.append("No title was supplied")

    doc.save(out, garbage=3, deflate=True, clean=True)
    doc.close()

    try:
        with pikepdf.open(out, allow_overwriting_input=True) as pdf:
            if "lang" in rules:
                lang = p.get("language") or "en-US"
                pdf.Root["/Lang"] = pikepdf.String(lang)
                applied.append(f"Set document language to {lang}")

            if "displayTitle" in rules:
                prefs = pdf.Root.get("/ViewerPreferences")
                if prefs is None:
                    pdf.Root["/ViewerPreferences"] = pdf.make_indirect(pikepdf.Dictionary())
                    prefs = pdf.Root["/ViewerPreferences"]
                prefs["/DisplayDocTitle"] = True
                applied.append("Readers will now show the title instead of the file name")

            if "marked" in rules:
                if "/StructTreeRoot" in pdf.Root:
                    mark = pdf.Root.get("/MarkInfo")
                    if mark is None:
                        pdf.Root["/MarkInfo"] = pdf.make_indirect(pikepdf.Dictionary())
                        mark = pdf.Root["/MarkInfo"]
                    mark["/Marked"] = True
                    applied.append("Declared the document tagged")
                else:
                    skipped.append("Cannot declare the file tagged — it has no structure tree")

            pdf.save(out, linearize=False)
    except Exception as e:
        skipped.append(f"Structural edits failed: {e}")

    result(output=out, applied=applied, skipped=skipped)


def cmd_reading_text(p):
    """Flow the document out as clean reading text, in reading order, for the
    reflow view and read-aloud."""
    doc = open_doc(p["input"], p.get("password"))
    pages = parse_pages(p.get("pages"), doc.page_count)
    out = []
    for i, pno in enumerate(pages):
        progress(i / max(1, len(pages)), f"Reading page {pno + 1}")
        page = doc[pno]
        blocks = []
        for b in page.get_text("dict").get("blocks", []):
            if b.get("type") != 0:
                continue
            sizes, text = [], []
            for line in b.get("lines", []):
                parts = [sp.get("text", "") for sp in line.get("spans", [])]
                for sp in line.get("spans", []):
                    sizes.append(sp.get("size", 11))
                text.append("".join(parts))
            body = " ".join(t.strip() for t in text if t.strip())
            if not body:
                continue
            avg = sum(sizes) / len(sizes) if sizes else 11
            blocks.append({"text": body, "size": round(avg, 1),
                           "bbox": [round(v, 1) for v in b["bbox"]]})
        # sort roughly into reading order: top to bottom, then left to right
        blocks.sort(key=lambda b: (round(b["bbox"][1] / 8), b["bbox"][0]))
        body_size = 11.0
        if blocks:
            counts = {}
            for b in blocks:
                counts[b["size"]] = counts.get(b["size"], 0) + len(b["text"])
            body_size = max(counts, key=counts.get)
        for b in blocks:
            b["heading"] = b["size"] >= body_size * 1.25
        out.append({"page": pno, "blocks": blocks})
    result(pages=out)


# --------------------------------------------------------------------------
# Acrobat-style extras
# --------------------------------------------------------------------------

def cmd_header_footer(p):
    """Headers and footers with left/centre/right fields, plus Bates numbering."""
    doc = open_doc(p["input"], p.get("password"))
    pages = parse_pages(p.get("pages"), doc.page_count)
    size = float(p.get("fontSize", 9))
    margin = float(p.get("margin", 28))
    color = p.get("color", [0.25, 0.25, 0.28])
    font = pick_font(p.get("font"), p.get("bold"), p.get("italic"))
    start = int(p.get("startAt", 1))
    bates_prefix = p.get("batesPrefix", "")
    bates_digits = int(p.get("batesDigits", 6))
    bates_start = int(p.get("batesStart", 1))
    stamp_date = p.get("date") or ""
    stamp_time = p.get("time") or ""
    name = os.path.splitext(os.path.basename(p["input"]))[0]

    slots = [("header", "left"), ("header", "center"), ("header", "right"),
             ("footer", "left"), ("footer", "center"), ("footer", "right")]

    for i, pno in enumerate(pages):
        progress(i / max(1, len(pages)), f"Stamping page {pno + 1}")
        page = doc[pno]
        rect = page.rect
        n = start + i
        bates = f"{bates_prefix}{str(bates_start + i).zfill(bates_digits)}"
        for band, align in slots:
            template = (p.get(band) or {}).get(align, "")
            if not template:
                continue
            label = (template.replace("{n}", str(n))
                             .replace("{total}", str(len(pages)))
                             .replace("{page}", str(pno + 1))
                             .replace("{bates}", bates)
                             .replace("{filename}", name)
                             .replace("{date}", stamp_date)
                             .replace("{time}", stamp_time))
            width = pymupdf.get_text_length(label, fontname=font, fontsize=size)
            if align == "left":
                x = margin
            elif align == "right":
                x = rect.width - margin - width
            else:
                x = (rect.width - width) / 2
            y = margin if band == "header" else rect.height - margin - size
            page.insert_textbox(pymupdf.Rect(x, y, x + width + 2, y + size * 1.5),
                                label, fontname=font, fontsize=size,
                                color=color, overlay=True)
    save_optimised(doc, p["output"], linear=False)
    result(output=p["output"], pagesStamped=len(pages))


def cmd_tables_to_excel(p):
    """Pull tables out into a real spreadsheet, one sheet per page."""
    import openpyxl
    from openpyxl.styles import Font as XLFont, PatternFill

    doc = open_doc(p["input"], p.get("password"))
    pages = parse_pages(p.get("pages"), doc.page_count)
    book = openpyxl.Workbook()
    book.remove(book.active)
    found = 0

    def usable(rows, strict):
        """Reject the noise the whitespace strategy invents out of prose."""
        if len(rows) < 2 or max((len(r) for r in rows), default=0) < 2:
            return False
        cells = [str(c).strip() for r in rows for c in r if c and str(c).strip()]
        if len(cells) < max(4, len(rows)):
            return False
        if not strict:
            return True
        # A real unruled table has short, consistently shaped cells. Running
        # prose chopped into columns does not.
        if len(rows) < 3:
            return False
        lengths = sorted(len(c) for c in cells)
        median = lengths[len(lengths) // 2]
        if median > 22 or max(lengths) > 90:
            return False
        counts = [sum(1 for c in r if c and str(c).strip()) for r in rows]
        common = max(set(counts), key=counts.count)
        if common < 2 or counts.count(common) < len(counts) * 0.6:
            return False
        # cells that read like sentences are a giveaway
        sentences = sum(1 for c in cells if c.count(" ") >= 4)
        return sentences <= len(cells) * 0.25

    # Unruled detection is opt-in: it is guesswork, and on prose it guesses badly.
    strategies = ["lines", "lines_strict"]
    if p.get("includeUnruled"):
        strategies.append("text")

    for i, pno in enumerate(pages):
        progress(i / max(1, len(pages)), f"Looking for tables on page {pno + 1}")
        page = doc[pno]
        found_here = []
        # ruled tables first; fall back to whitespace alignment only if needed
        for strategy in strategies:
            try:
                located = page.find_tables(strategy=strategy)
            except Exception:
                continue
            candidates = located.tables if hasattr(located, "tables") else list(located)
            good = []
            for table in candidates:
                try:
                    rows = table.extract()
                except Exception:
                    continue
                if usable(rows, strict=(strategy == "text")):
                    good.append(rows)
            if good:
                found_here = good
                break

        for t_index, rows in enumerate(found_here):
            if not rows:
                continue
            found += 1
            title = f"Page {pno + 1}" + (f" ({t_index + 1})" if t_index else "")
            sheet = book.create_sheet(title[:31])
            for r, row in enumerate(rows, start=1):
                for c, value in enumerate(row, start=1):
                    sheet.cell(row=r, column=c,
                               value=("" if value is None else str(value).replace("\n", " ")))
            for cell in sheet[1]:
                cell.font = XLFont(bold=True)
                cell.fill = PatternFill("solid", fgColor="EFEFF4")
            for column in sheet.columns:
                longest = max((len(str(c.value or "")) for c in column), default=8)
                sheet.column_dimensions[column[0].column_letter].width = min(60, max(9, longest + 2))

    if not found:
        raise EngineError(
            "No tables were detected. RiftPDF looks for tables with ruled lines; "
            "turn on “Include tables without borders” to also guess at "
            "whitespace-aligned ones.")

    book.save(p["output"])
    result(output=p["output"], tables=found, sheets=len(book.sheetnames))


def cmd_audit_space(p):
    """Where the bytes actually went — Acrobat's 'audit space usage'."""
    src = p["input"]
    total = os.path.getsize(src)
    buckets = {"Images": 0, "Fonts": 0, "Content streams": 0,
               "Embedded files": 0, "Metadata": 0, "Other": 0}
    with pikepdf.open(src, password=p.get("password") or "") as pdf:
        count = len(pdf.objects)
        for i, obj in enumerate(pdf.objects):
            if i % 500 == 0:
                progress(i / max(1, count), "Measuring objects")
            try:
                if not isinstance(obj, pikepdf.Stream):
                    continue
                raw = obj.read_raw_bytes()
                size = len(raw)
                subtype = str(obj.get("/Subtype", ""))
                otype = str(obj.get("/Type", ""))
                if subtype == "/Image":
                    buckets["Images"] += size
                elif "/FontFile" in obj.keys() or otype == "/Font" or subtype in ("/Type1C", "/CIDFontType0C", "/OpenType"):
                    buckets["Fonts"] += size
                elif subtype == "/XML" or otype == "/Metadata":
                    buckets["Metadata"] += size
                elif otype == "/EmbeddedFile":
                    buckets["Embedded files"] += size
                elif otype == "/XObject" or obj.get("/Length") is not None:
                    buckets["Content streams"] += size
                else:
                    buckets["Other"] += size
            except Exception:
                continue

    accounted = sum(buckets.values())
    buckets["Structure and overhead"] = max(0, total - accounted)
    rows = [{"category": k, "bytes": v, "human": human(v),
             "percent": round(100.0 * v / total, 1) if total else 0}
            for k, v in sorted(buckets.items(), key=lambda kv: -kv[1]) if v > 0]
    result(total=total, totalHuman=human(total), categories=rows)


def _tessdata_dir():
    """Tesseract keeps its language data outside the binary; find it."""
    env = os.environ.get("TESSDATA_PREFIX")
    if env and os.path.isdir(env):
        return env
    for candidate in ("/opt/homebrew/share/tessdata", "/usr/local/share/tessdata",
                      "/usr/share/tesseract-ocr/5/tessdata",
                      "/usr/share/tesseract-ocr/4.00/tessdata",
                      "/usr/share/tessdata",
                      r"C:\\Program Files\\Tesseract-OCR\\tessdata",
                      r"C:\\Program Files (x86)\\Tesseract-OCR\\tessdata"):
        if os.path.isdir(candidate):
            return candidate
    exe = which("tesseract")
    if exe:
        guess = os.path.join(os.path.dirname(os.path.dirname(exe)), "share", "tessdata")
        if os.path.isdir(guess):
            return guess
    return None


def cmd_ocr(p):
    """Recognise text on scanned pages and weld in a searchable text layer.

    Pages that already carry real text are copied through untouched, so a mixed
    document does not get flattened into images.
    """
    if not which("tesseract"):
        raise EngineError(
            "Tesseract is not installed. On Windows install it from "
            "https://github.com/UB-Mannheim/tesseract/wiki, on macOS run "
            "'brew install tesseract'.")
    tessdata = _tessdata_dir()
    if not tessdata:
        raise EngineError("Tesseract is installed but its language data was not found. "
                          "Set TESSDATA_PREFIX to the tessdata folder.")
    os.environ["TESSDATA_PREFIX"] = tessdata

    doc = open_doc(p["input"], p.get("password"))
    pages = parse_pages(p.get("pages"), doc.page_count)
    language = p.get("language", "eng")
    dpi = int(p.get("dpi", 300))
    force = bool(p.get("force"))

    out = pymupdf.open()
    recognised = 0
    for i, pno in enumerate(pages):
        progress(i / max(1, len(pages)), f"Reading page {pno + 1} of {len(pages)}")
        page = doc[pno]
        existing = page.get_text("text").strip()
        if len(existing) > 20 and not force:
            out.insert_pdf(doc, from_page=pno, to_page=pno)
            continue
        pix = page.get_pixmap(dpi=dpi)
        try:
            data = pix.pdfocr_tobytes(language=language, tessdata=tessdata)
        except Exception as exc:
            raise EngineError(f"Tesseract failed on page {pno + 1}: {exc}")
        piece = pymupdf.open("pdf", data)
        out.insert_pdf(piece)
        piece.close()
        recognised += 1

    progress(0.95, "Writing")
    save_optimised(out, p["output"], linear=False)
    characters = 0
    check = pymupdf.open(p["output"])
    for page in check:
        characters += len(page.get_text("text").strip())
    check.close()
    result(output=p["output"], pagesRecognised=recognised,
           pagesCopied=len(pages) - recognised, characters=characters)


def cmd_merge_annotations(p):
    """Put the markup from one file onto the well-compressed bytes of another.

    PDFKit re-encodes images when it writes a document, which can more than
    double a scanned page. So the app keeps the engine's output and, on save,
    transplants only the annotations onto it rather than accepting PDFKit's
    whole rewrite.
    """
    base, overlay, out = p["base"], p["overlay"], p["output"]

    with pikepdf.open(overlay) as source:
        if len(source.pages) == 0:
            raise EngineError("The overlay document has no pages.")
        # Form fields live in both the page annotations and the AcroForm field
        # tree; splitting them across two files would break the form, so leave
        # these documents to PDFKit.
        if "/AcroForm" in source.Root:
            raise EngineError("ACROFORM_PRESENT")
        page_count = len(source.pages)

    with pikepdf.open(base) as target:
        if len(target.pages) != page_count:
            raise EngineError("PAGE_COUNT_MISMATCH")

    with pikepdf.open(base) as target, pikepdf.open(overlay) as source:
        moved = 0
        for tpage, spage in zip(target.pages, source.pages):
            annots = spage.get("/Annots")
            if annots is None or len(annots) == 0:
                if "/Annots" in tpage:
                    del tpage["/Annots"]
                continue
            copied = [target.copy_foreign(a) for a in annots]
            tpage["/Annots"] = target.make_indirect(pikepdf.Array(copied))
            moved += len(copied)
        target.save(out, linearize=False)

    # never hand back something that lost markup
    check = pymupdf.open(out)
    found = sum(len(list(page.annots())) for page in check)
    pages = check.page_count
    check.close()
    if found != moved:
        raise EngineError("Annotations did not survive the transplant.")

    result(output=out, annotations=moved, pageCount=pages,
           size=os.path.getsize(out))


# --------------------------------------------------------------------------
# dispatch
# --------------------------------------------------------------------------

COMMANDS = {
    "info": cmd_info,
    "compress": cmd_compress,
    "compress_target": cmd_compress_target,
    "metadata_set": cmd_metadata_set,
    "metadata_strip": cmd_metadata_strip,
    "sanitize": cmd_sanitize,
    "merge": cmd_merge,
    "split": cmd_split,
    "images_to_pdf": cmd_images_to_pdf,
    "pdf_to_images": cmd_pdf_to_images,
    "extract_images": cmd_extract_images,
    "pdf_to_word": cmd_pdf_to_word,
    "pdf_to_text": cmd_pdf_to_text,
    "office_to_pdf": cmd_office_to_pdf,
    "watermark": cmd_watermark,
    "page_numbers": cmd_page_numbers,
    "redact": cmd_redact,
    "redact_search": cmd_redact_search,
    "text_spans": cmd_text_spans,
    "text_edit": cmd_text_edit,
    "add_text": cmd_add_text,
    "place_images": cmd_place_images,
    "ocr_layer": cmd_ocr_layer,
    "encrypt": cmd_encrypt,
    "decrypt": cmd_decrypt,
    "page_ops": cmd_page_ops,
    "bookmarks_set": cmd_bookmarks_set,
    "attachments": cmd_attachments,
    "form_fields": cmd_form_fields,
    "form_fill": cmd_form_fill,
    "flatten": cmd_flatten,
    "repair": cmd_repair,
    "linearize": cmd_linearize,
    "compare": cmd_compare,
    "accessibility_check": cmd_accessibility_check,
    "accessibility_fix": cmd_accessibility_fix,
    "reading_text": cmd_reading_text,
    "header_footer": cmd_header_footer,
    "tables_to_excel": cmd_tables_to_excel,
    "audit_space": cmd_audit_space,
    "merge_annotations": cmd_merge_annotations,
    "ocr": cmd_ocr,
}


def main():
    if len(sys.argv) < 2:
        emit({"type": "error", "message": "No command given",
              "detail": "commands: " + ", ".join(sorted(COMMANDS))})
        return 2
    name = sys.argv[1]
    if name == "--selftest":
        payload = {"type": "result"}
        payload.update(selftest())
        emit(payload)
        return 0
    fn = COMMANDS.get(name)
    if not fn:
        emit({"type": "error", "message": f"Unknown command '{name}'"})
        return 2
    try:
        if len(sys.argv) > 2 and os.path.exists(sys.argv[2]):
            with open(sys.argv[2], "r", encoding="utf-8") as fh:
                payload = json.load(fh)
        else:
            payload = json.loads(sys.argv[2]) if len(sys.argv) > 2 else {}
    except Exception as e:
        emit({"type": "error", "message": f"Bad payload: {e}"})
        return 2
    try:
        fn(payload)
        return 0
    except EngineError as e:
        emit({"type": "error", "message": str(e)})
        return 1
    except Exception as e:
        emit({"type": "error", "message": str(e), "detail": traceback.format_exc()})
        return 1


if __name__ == "__main__":
    sys.exit(main())
